"""Bounded import of action-plan rows from a spreadsheet or UTF-8 CSV."""
import csv
from io import BytesIO, StringIO
from itertools import islice
from zipfile import BadZipFile, ZipFile

from rest_framework.exceptions import ValidationError


MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 200
HEADERS = ('building_code', 'expert_phone_number', 'elevator_ids')
ALIASES = {'outlet_code': 'building_code', 'promoter_phone_number': 'expert_phone_number'}


def _header(value):
    return ALIASES.get(str(value or '').strip().lower(), str(value or '').strip().lower())


def _rows(upload):
    name = (upload.name or '').lower()
    content = upload.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValidationError({'file': 'حجم فایل باید حداکثر ۲ مگابایت باشد.'})
    if name.endswith('.csv'):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeError:
            raise ValidationError({'file': 'فایل CSV باید با UTF-8 ذخیره شود.'})
        return islice(csv.reader(StringIO(text)), MAX_ROWS + 2)
    if name.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
            with ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if (len(members) > 100 or
                        sum(member.file_size for member in members) > 10 * 1024 * 1024 or
                        any(member.file_size > 8 * 1024 * 1024 for member in members)):
                    raise ValidationError({'file': 'محتوای فایل XLSX بیش از حد بزرگ است.'})
            workbook = load_workbook(BytesIO(content), read_only=True,
                                     data_only=True, keep_links=False)
            try:
                return list(islice(workbook.active.iter_rows(max_col=20, values_only=True),
                                   MAX_ROWS + 2))
            finally:
                workbook.close()
        except (BadZipFile, ValueError, KeyError, OSError):
            raise ValidationError({'file': 'فایل XLSX معتبر نیست.'})
    raise ValidationError({'file': 'فقط فایل XLSX یا CSV پذیرفته می‌شود.'})


def parse_action_plan_rows(upload):
    rows = iter(_rows(upload))
    header = next(rows, None)
    if not header:
        raise ValidationError({'file': 'فایل خالی است.'})
    columns = [_header(value) for value in header]
    if any(columns.count(name) != 1 for name in HEADERS[:2]) or len(set(columns)) != len(columns):
        raise ValidationError({'file': 'ستون‌های building_code و expert_phone_number باید یکتا باشند.'})
    positions = {name: columns.index(name) for name in HEADERS if name in columns}
    result = []
    for source_row, values in enumerate(rows, 2):
        if source_row > MAX_ROWS + 1:
            raise ValidationError({'file': f'حداکثر {MAX_ROWS} ردیف داده پذیرفته می‌شود.'})
        if not any(value is not None and str(value).strip() for value in values):
            continue
        def cell(name):
            index = positions.get(name)
            return values[index] if index is not None and index < len(values) else None
        code = str(cell('building_code') or '').strip()
        phone_value = cell('expert_phone_number')
        phone = str(phone_value or '').strip()
        if phone.endswith('.0') and phone[:-2].isdigit():
            phone = phone[:-2]
        if phone.isdigit() and len(phone) == 10:
            phone = '0' + phone
        if not code or len(code) > 25 or not phone or len(phone) > 32:
            raise ValidationError({'file': f'کد ساختمان یا شماره کارشناس در ردیف {source_row} معتبر نیست.'})
        raw_ids = str(cell('elevator_ids') or '').strip()
        parts = [part.strip() for part in raw_ids.replace('؛', ',').replace(';', ',').split(',') if part.strip()]
        if any(not part.isascii() or not part.isdigit() for part in parts) or len(parts) > 20:
            raise ValidationError({'file': f'شناسه آسانسور در ردیف {source_row} معتبر نیست.'})
        ids = [int(part) for part in parts]
        if len(ids) != len(set(ids)) or any(value <= 0 for value in ids):
            raise ValidationError({'file': f'شناسه آسانسور در ردیف {source_row} تکراری یا نامعتبر است.'})
        result.append({'building_code': code, 'expert_phone_number': phone,
                       'elevator_ids': ids, 'source_row': source_row})
    if not result:
        raise ValidationError({'file': 'فایل هیچ ردیف برنامه‌ای ندارد.'})
    return result
